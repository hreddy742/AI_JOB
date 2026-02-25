"""Resume editing + version history service."""

from __future__ import annotations

import json
from difflib import SequenceMatcher
from io import BytesIO
from uuid import UUID

import httpx
from docx import Document
from fastapi import HTTPException, status
from minio import Minio
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from core.config import settings
from core.dependencies import apply_tenant_rls, apply_user_rls
from db.models.resume import Resume
from db.models.resume_version import ResumeVersion


def _minio_client() -> Minio:
    secure = settings.MINIO_ENDPOINT.startswith("https://")
    endpoint = settings.MINIO_ENDPOINT.replace("https://", "").replace("http://", "")
    return Minio(endpoint, access_key=settings.MINIO_ACCESS_KEY, secret_key=settings.MINIO_SECRET_KEY, secure=secure)


async def save_version(
    db: AsyncSession,
    resume_id: UUID,
    user_id: UUID,
    content_text: str,
    change_summary: str,
    content_parsed: dict | None = None,
) -> ResumeVersion:
    """Persist a resume version snapshot and cap history to 50."""

    max_version = (
        await db.execute(select(func.max(ResumeVersion.version_number)).where(ResumeVersion.resume_id == resume_id))
    ).scalar_one_or_none()
    next_version = int(max_version or 0) + 1
    version = ResumeVersion(
        user_id=user_id,
        resume_id=resume_id,
        version_number=next_version,
        content_text=content_text,
        content_parsed=content_parsed,
        change_summary=change_summary[:500],
    )
    db.add(version)
    await db.flush()

    old_rows = (
        await db.execute(
            select(ResumeVersion)
            .where(ResumeVersion.resume_id == resume_id)
            .order_by(ResumeVersion.version_number.desc())
            .offset(50)
        )
    ).scalars().all()
    for row in old_rows:
        await db.delete(row)
    await db.commit()
    await db.refresh(version)
    return version


def _similarity(a: str, b: str) -> float:
    return SequenceMatcher(None, a.strip().lower(), b.strip().lower()).ratio()


async def edit_docx_resume(
    db: AsyncSession,
    user_id: UUID,
    tenant_id: UUID,
    resume_id: UUID,
    original_text: str,
    new_text: str,
    change_summary: str,
) -> Resume:
    """Replace one DOCX paragraph by fuzzy match while preserving style."""

    await apply_tenant_rls(db, tenant_id)
    await apply_user_rls(db, user_id)
    resume = (
        await db.execute(select(Resume).where(Resume.id == resume_id, Resume.user_id == user_id, Resume.tenant_id == tenant_id))
    ).scalar_one_or_none()
    if resume is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Resume not found")
    if resume.file_type != "docx":
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="DOCX edit endpoint supports only docx")

    minio = _minio_client()
    data = minio.get_object(settings.MINIO_BUCKET, resume.file_path).read()
    doc = Document(BytesIO(data))

    best_idx = -1
    best_score = 0.0
    for idx, para in enumerate(doc.paragraphs):
        score = _similarity(para.text, original_text)
        if score > best_score:
            best_score = score
            best_idx = idx
    if best_idx < 0 or best_score < 0.7:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Could not locate paragraph to edit")

    target = doc.paragraphs[best_idx]
    runs = list(target.runs)
    if runs:
        first = runs[0]
        first.text = new_text
        for run in runs[1:]:
            run.text = ""
    else:
        target.text = new_text

    out = BytesIO()
    doc.save(out)
    out.seek(0)
    minio.put_object(
        settings.MINIO_BUCKET,
        resume.file_path,
        out,
        length=len(out.getvalue()),
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    resume.raw_text = (resume.raw_text or "").replace(original_text, new_text, 1)
    resume.original_text = resume.raw_text
    await save_version(db, resume.id, user_id, resume.original_text, change_summary)
    await db.commit()
    await db.refresh(resume)
    return resume


async def improve_bullet(original_bullet: str, tense: str, title: str, company: str, target_role: str) -> str:
    """Generate one grounded bullet rewrite suggestion."""

    prompt = f"""
Improve this single resume bullet point.
Rules: no invented facts, metrics, companies, or scope.
Tense: {tense}
Original bullet: {original_bullet}
Context: {title} at {company}
Target role: {target_role}
Respond with ONLY improved bullet text.
""".strip()
    async with httpx.AsyncClient(timeout=120.0) as client:
        response = await client.post(
            f"{settings.OLLAMA_BASE_URL}/api/chat",
            json={
                "model": settings.WRITER_MODEL,
                "messages": [{"role": "system", "content": prompt}],
                "stream": False,
                "options": {"temperature": 0.2},
            },
        )
        response.raise_for_status()
        return response.json().get("message", {}).get("content", "").strip()


async def restore_resume_version(db: AsyncSession, user_id: UUID, tenant_id: UUID, resume_id: UUID, version_number: int) -> Resume:
    """Restore stored text snapshot to current resume."""

    await apply_tenant_rls(db, tenant_id)
    await apply_user_rls(db, user_id)
    version = (
        await db.execute(
            select(ResumeVersion).where(
                ResumeVersion.resume_id == resume_id, ResumeVersion.user_id == user_id, ResumeVersion.version_number == version_number
            )
        )
    ).scalar_one_or_none()
    if version is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Version not found")
    resume = (
        await db.execute(select(Resume).where(Resume.id == resume_id, Resume.user_id == user_id, Resume.tenant_id == tenant_id))
    ).scalar_one_or_none()
    if resume is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Resume not found")
    resume.original_text = version.content_text
    resume.raw_text = version.content_text
    await save_version(db, resume_id, user_id, version.content_text, f"restored from version {version_number}")
    await db.commit()
    await db.refresh(resume)
    return resume
