"""Resume upload and parsing pipeline."""

from __future__ import annotations

import json
import re
import unicodedata
from collections import Counter
from datetime import UTC, datetime
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any
from uuid import UUID, uuid4

import chromadb
import chardet
import httpx
import pytesseract
from arq import create_pool
from arq.connections import RedisSettings
from dateutil import parser as date_parser
from docx import Document
from fastapi import HTTPException, UploadFile, status
from minio import Minio
from pdf2image import convert_from_bytes
from pdfminer.high_level import extract_text
from pdfminer.layout import LAParams
from pydantic import BaseModel, Field, ValidationError
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from core.config import settings
from core.dependencies import apply_tenant_rls, apply_user_rls
from db.models.resume import Resume
from db.models.resume_parsed_data import ResumeParsedData
from services.embedding_service import store_resume_embedding

MAX_SIZE_BYTES = 10 * 1024 * 1024
MAX_RAW_TEXT_LENGTH = 50_000

SECTION_PATTERNS: dict[str, list[str]] = {
    "summary": [r"professional summary", r"summary", r"objective", r"about me", r"profile", r"career overview"],
    "experience": [r"work experience", r"experience", r"employment history", r"professional background", r"work history"],
    "education": [r"education", r"academic background", r"degrees", r"academic qualifications"],
    "skills": [r"skills", r"technical skills", r"competencies", r"expertise", r"technologies", r"proficiencies"],
    "projects": [r"projects", r"personal projects", r"portfolio", r"notable projects", r"key projects"],
    "certifications": [r"certifications", r"licenses", r"credentials", r"professional development"],
    "awards": [r"awards", r"honors", r"achievements", r"recognition"],
    "publications": [r"publications", r"papers", r"research"],
    "volunteer": [r"volunteer", r"community service", r"non-profit"],
    "languages": [r"languages", r"spoken languages", r"linguistic"],
}

STOP_WORDS = {
    "the",
    "and",
    "for",
    "with",
    "to",
    "of",
    "in",
    "on",
    "a",
    "an",
    "at",
    "by",
    "as",
    "from",
    "using",
    "used",
    "that",
    "this",
    "was",
    "were",
    "is",
    "are",
}


class WorkEntry(BaseModel):
    company: str | None = None
    title: str | None = None
    location: str | None = None
    start_date: str | None = None
    end_date: str | None = None
    is_current: bool = False
    bullets: list[str] = Field(default_factory=list)
    technologies: list[str] = Field(default_factory=list)
    achievements: list[str] = Field(default_factory=list)
    raw_text: str | None = None


class EducationEntry(BaseModel):
    institution: str | None = None
    degree: str | None = None
    field_of_study: str | None = None
    start_date: str | None = None
    end_date: str | None = None
    gpa: str | None = None
    honors: str | None = None
    relevant_courses: list[str] = Field(default_factory=list)
    activities: list[str] = Field(default_factory=list)


class SkillOutput(BaseModel):
    technical: list[str] = Field(default_factory=list)
    soft: list[str] = Field(default_factory=list)
    languages: list[str] = Field(default_factory=list)
    certifications: list[str] = Field(default_factory=list)


def get_minio_client() -> Minio:
    """Create MinIO client for resume storage."""

    secure = settings.MINIO_ENDPOINT.startswith("https://")
    endpoint = settings.MINIO_ENDPOINT.replace("https://", "").replace("http://", "")
    return Minio(
        endpoint,
        access_key=settings.MINIO_ACCESS_KEY,
        secret_key=settings.MINIO_SECRET_KEY,
        secure=secure,
    )


def _normalize_text(text: str) -> str:
    text = unicodedata.normalize("NFKC", text)
    text = text.replace("\u2014", "-").replace("\u2013", "-").replace("\u2018", "'").replace("\u2019", "'")
    text = text.replace("\u201c", '"').replace("\u201d", '"')
    text = re.sub(r"^[\u2022\u25e6\u25c6\-\*]\s*", "- ", text, flags=re.MULTILINE)
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = "\n".join(line.rstrip() for line in text.splitlines())
    return text.strip()


def _remove_repeated_short_lines_by_page(text: str) -> str:
    pages = [p for p in re.split(r"\f", text) if p.strip()]
    if len(pages) <= 1:
        return text
    top_lines: list[str] = []
    bottom_lines: list[str] = []
    for page in pages:
        lines = [l.strip() for l in page.splitlines() if l.strip()]
        if not lines:
            continue
        top_lines.extend([l for l in lines[:2] if len(l) < 50])
        bottom_lines.extend([l for l in lines[-2:] if len(l) < 50])
    repeated = {line for line, count in Counter(top_lines + bottom_lines).items() if count >= 2}
    cleaned_pages: list[str] = []
    for page in pages:
        lines = [l for l in page.splitlines() if l.strip() not in repeated]
        cleaned_pages.append("\n".join(lines))
    return "\n\n".join(cleaned_pages)


async def extract_raw_text(file_bytes: bytes, file_type: str) -> str:
    """Extract normalized text from supported resume file types."""

    text = ""
    normalized_type = file_type.lower()
    if normalized_type == "pdf":
        laparams = LAParams(line_overlap=0.5, char_margin=2.0, line_margin=0.5)
        text = extract_text(BytesIO(file_bytes), laparams=laparams)
        text = _remove_repeated_short_lines_by_page(text)
        if len(text.strip()) < 100:
            pages = convert_from_bytes(file_bytes, dpi=300)
            text = "\n\n".join(pytesseract.image_to_string(page) for page in pages)
    elif normalized_type == "docx":
        doc = Document(BytesIO(file_bytes))
        parts: list[str] = []
        for para in doc.paragraphs:
            style = (para.style.name or "").lower()
            prefix = ""
            if style.startswith("heading"):
                prefix = "## "
            elif "list bullet" in style or "list number" in style:
                prefix = "- "
            if para.text.strip():
                parts.append(f"{prefix}{para.text}")
        for table in doc.tables:
            for row in table.rows:
                row_text = "\t".join(cell.text.strip() for cell in row.cells)
                if row_text.strip():
                    parts.append(row_text)
        text = "\n".join(parts)
    else:
        detected = chardet.detect(file_bytes) or {}
        encoding = detected.get("encoding") or "utf-8"
        try:
            text = file_bytes.decode(encoding)
        except Exception:
            text = file_bytes.decode("utf-8", errors="replace")

    text = _normalize_text(text)
    if len(text) < 100:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Could not extract text from resume")
    if len(text) > MAX_RAW_TEXT_LENGTH:
        text = text[:MAX_RAW_TEXT_LENGTH]
    if not (
        re.search(r"[\w\.-]+@[\w\.-]+\.\w{2,}", text)
        or re.search(r"\b(?:19|20)\d{2}\b", text)
        or re.search(r"\d{3}[\-\.\s]\d{3}[\-\.\s]\d{4}", text)
    ):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Resume text quality check failed")
    return text


def detect_sections(text: str) -> dict[str, str]:
    """Detect major resume sections using rule-based header matching."""

    lines = text.splitlines()
    section_lines: dict[str, list[str]] = {}
    current_section = "general"
    section_lines[current_section] = []
    for line in lines:
        stripped = line.strip()
        if stripped and len(stripped) <= 60:
            lowered = stripped.lower()
            matched = None
            for section, patterns in SECTION_PATTERNS.items():
                if any(re.fullmatch(pattern, lowered, flags=re.IGNORECASE) for pattern in patterns):
                    matched = section
                    break
            if matched:
                current_section = matched
                section_lines.setdefault(current_section, [])
                continue
        section_lines.setdefault(current_section, []).append(line)
    return {key: "\n".join(value).strip() for key, value in section_lines.items() if "\n".join(value).strip()}


def extract_contact_regex(text: str) -> dict[str, Any]:
    """Extract contact block with deterministic regex and top-line heuristics."""

    lines = [line.strip() for line in text.splitlines() if line.strip()]
    top_line = lines[0] if lines else ""
    email_regex = r"[\w\.-]+@[\w\.-]+\.\w{2,}"
    email_spans = [m.span() for m in re.finditer(email_regex, text)]

    def _overlaps_email(start: int, end: int) -> bool:
        for email_start, email_end in email_spans:
            if start < email_end and end > email_start:
                return True
        return False

    website = None
    for match in re.finditer(
        r"(?:https?://)?(?:www\.)?[A-Za-z0-9\-]+\.[A-Za-z]{2,}(?:/[^\s]*)?(?=$|[/?#\s])",
        text,
    ):
        candidate = match.group(0)
        if _overlaps_email(match.start(), match.end()):
            continue
        lowered = candidate.lower()
        if "linkedin.com" in lowered or "github.com" in lowered:
            continue
        website = candidate
        break

    return {
        "name": top_line if top_line and len(top_line.split()) <= 6 else None,
        "email": next(iter(re.findall(email_regex, text)), None),
        "phone": next(
            iter(re.findall(r"(\+?1?[\s\-\.]?\(?\d{3}\)?[\s\-\.]*\d{3}[\s\-\.]*\d{4})", text)),
            None,
        ),
        "linkedin": next(iter(re.findall(r"(?:https?://)?(?:www\.)?linkedin\.com/in/[\w\-]+", text, flags=re.I)), None),
        "github": next(iter(re.findall(r"(?:https?://)?(?:www\.)?github\.com/[\w\-]+", text, flags=re.I)), None),
        "website": website,
    }


async def _call_ollama_json(system_prompt: str, user_content: str, model: str, temperature: float = 0.0) -> dict[str, Any]:
    async with httpx.AsyncClient(timeout=120.0) as client:
        resp = await client.post(
            f"{settings.OLLAMA_BASE_URL}/api/chat",
            json={
                "model": model,
                "messages": [
                    {"role": "system", "content": system_prompt},
                    {"role": "user", "content": user_content},
                ],
                "format": "json",
                "stream": False,
                "options": {"temperature": temperature},
            },
        )
        resp.raise_for_status()
        payload = resp.json()
    content = payload.get("message", {}).get("content", "{}")
    return content if isinstance(content, dict) else json.loads(content)


async def _parse_json_with_retry(system_prompt: str, user_content: str, model: str) -> dict[str, Any]:
    try:
        return await _call_ollama_json(system_prompt, user_content, model, 0.0)
    except Exception:
        retry_prompt = f"{system_prompt}\nIMPORTANT: Your previous response was not valid JSON. Output ONLY the JSON."
        return await _call_ollama_json(retry_prompt, user_content, model, 0.0)


def _parse_dt(value: str | None) -> datetime | None:
    if not value:
        return None
    try:
        dt = date_parser.parse(value, fuzzy=True)
        if dt.tzinfo is None:
            return dt.replace(tzinfo=UTC)
        return dt.astimezone(UTC)
    except Exception:
        return None


def compute_metadata(parsed: dict[str, Any]) -> dict[str, Any]:
    """Compute deterministic metadata from parsed resume payload."""

    work = parsed.get("work_experience", []) or []
    total_months = 0
    titles: list[str] = []
    all_bullets: list[str] = []
    for item in work:
        titles.append((item.get("title") or "").strip())
        all_bullets.extend(item.get("bullets") or [])
        start = _parse_dt(item.get("start_date"))
        end_raw = (item.get("end_date") or "").lower()
        end = datetime.now(UTC) if ("present" in end_raw or "current" in end_raw) else _parse_dt(item.get("end_date"))
        if start and end and end >= start:
            total_months += max((end.year - start.year) * 12 + (end.month - start.month), 0)
    years = round(total_months / 12.0, 1)

    combined_titles = " ".join(t.lower() for t in titles)
    seniority = "entry"
    if any(k in combined_titles for k in ["director", "vp", "vice president", "head of"]):
        seniority = "executive"
    elif any(k in combined_titles for k in ["principal", "staff", "distinguished", "fellow"]):
        seniority = "lead"
    elif any(k in combined_titles for k in ["senior", "sr.", "lead", "manager"]) or years >= 7:
        seniority = "senior"
    elif years >= 3:
        seniority = "mid"

    skills_all = []
    for bucket in ["skills_technical", "skills_soft", "skills_languages", "skills_certifications"]:
        for skill in parsed.get(bucket, []) or []:
            if skill not in skills_all:
                skills_all.append(skill)

    scope = f"{combined_titles} {' '.join((parsed.get('skills_technical') or []))}".lower()
    domain = "other"
    domain_rules = [
        ("software_engineering", ["software engineer", "developer", "programmer", "backend", "frontend", "fullstack"]),
        ("data_science", ["data scientist", "ml engineer", "machine learning", "ai", "deep learning"]),
        ("data_engineering", ["data engineer", "spark", "kafka", "airflow", "dbt", "etl", "pipeline"]),
        ("devops", ["devops", "sre", "platform engineer", "infrastructure", "kubernetes"]),
        ("product", ["product manager", "product owner"]),
        ("design", ["ux designer", "ui designer", "product designer"]),
        ("security", ["security engineer", "cybersecurity", "penetration"]),
    ]
    for candidate, markers in domain_rules:
        if any(marker in scope for marker in markers):
            domain = candidate
            break

    tokens = re.findall(r"[A-Za-z][A-Za-z\+\#\-]{2,}", " ".join(all_bullets).lower())
    freq = Counter(t for t in tokens if t not in STOP_WORDS)
    industry_keywords = [word for word, _count in freq.most_common(30)]

    parsed["skills_all"] = skills_all
    parsed["total_years_experience"] = years
    parsed["seniority_level"] = seniority
    parsed["primary_domain"] = domain
    parsed["primary_job_titles"] = [title for title in titles if title][:3]
    parsed["industry_keywords"] = industry_keywords
    return parsed


def _infer_file_type(file_name: str, file_bytes: bytes) -> str:
    ext = Path(file_name).suffix.lower()
    if ext == ".pdf" and file_bytes[:4] == b"%PDF":
        return "pdf"
    if ext == ".docx" and file_bytes[:4] == b"PK\x03\x04":
        return "docx"
    if ext in {".txt", ".md"}:
        return "txt"
    try:
        file_bytes[:1024].decode("utf-8")
        return "txt"
    except Exception as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Unsupported resume file type") from exc


async def _queue_parse_job(resume_id: UUID) -> None:
    redis = await create_pool(RedisSettings.from_dsn(settings.REDIS_URL))
    try:
        await redis.enqueue_job("parse_resume", str(resume_id))
    finally:
        await redis.close()


async def upload_resume(
    file: UploadFile,
    user_id: UUID,
    tenant_id: UUID,
    label: str,
    db: AsyncSession,
    minio_client: Minio | None = None,
) -> Resume:
    """Validate, store, and enqueue resume parsing."""

    file_bytes = await file.read()
    if len(file_bytes) > MAX_SIZE_BYTES:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Resume exceeds 10MB limit")
    file_name = file.filename or f"resume-{uuid4()}.txt"
    file_type = _infer_file_type(file_name, file_bytes)
    object_id = uuid4()
    object_path = f"tenants/{tenant_id}/resumes/{user_id}/{object_id}.{file_type}"

    minio = minio_client or get_minio_client()
    if not minio.bucket_exists(settings.MINIO_BUCKET):
        minio.make_bucket(settings.MINIO_BUCKET)
    minio.put_object(
        settings.MINIO_BUCKET,
        object_path,
        BytesIO(file_bytes),
        length=len(file_bytes),
        content_type=file.content_type or "application/octet-stream",
    )

    await apply_tenant_rls(db, tenant_id)
    await apply_user_rls(db, user_id)
    resume = Resume(
        user_id=user_id,
        tenant_id=tenant_id,
        label=label or "My Resume",
        file_name=file_name,
        file_size_bytes=len(file_bytes),
        file_type=file_type,
        file_path=object_path,
        original_text="",
        raw_text=None,
        parse_status="pending",
    )
    db.add(resume)
    await db.commit()
    await db.refresh(resume)
    await _queue_parse_job(resume.id)
    return resume


async def parse_resume_pipeline(db: AsyncSession, resume_id: UUID) -> ResumeParsedData:
    """Run full parse pipeline for an existing resume row."""

    result = await db.execute(select(Resume).where(Resume.id == resume_id))
    resume = result.scalar_one_or_none()
    if resume is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Resume not found")

    await apply_tenant_rls(db, resume.tenant_id)
    await apply_user_rls(db, resume.user_id)
    resume.parse_status = "parsing"
    await db.commit()

    minio = get_minio_client()
    data = minio.get_object(settings.MINIO_BUCKET, resume.file_path).read()
    raw_text = await extract_raw_text(data, resume.file_type)
    sections = detect_sections(raw_text)
    contact = extract_contact_regex(raw_text)

    work_prompt = """
You are a resume parsing expert. Extract all work experience entries.
Copy content EXACTLY as written. Do not paraphrase, expand, or summarize.
Output ONLY this JSON:
{"work_experience": [...all entries in reverse chronological order...]}
""".strip()
    edu_prompt = """
Extract all education entries. Preserve all text exactly as written.
Output ONLY: {"education": [...entries...]}
""".strip()
    skill_prompt = """
Extract all skills. Categorize each. Do not add skills not in the text.
Output ONLY:
{"technical":[],"soft":[],"languages":[],"certifications":[]}
""".strip()
    projects_prompt = """
Extract projects from this section.
Output ONLY: {"projects":[{"name":"","description":"","technologies":[],"url":null,"bullets":[]}]}
""".strip()

    work_raw = await _parse_json_with_retry(work_prompt, sections.get("experience", ""), settings.WRITER_MODEL)
    edu_raw = await _parse_json_with_retry(edu_prompt, sections.get("education", ""), settings.WRITER_MODEL)
    skills_raw = await _parse_json_with_retry(skill_prompt, sections.get("skills", ""), settings.WRITER_MODEL)
    projects_raw = await _parse_json_with_retry(projects_prompt, sections.get("projects", ""), settings.WRITER_MODEL)

    try:
        work = [WorkEntry.model_validate(item).model_dump() for item in work_raw.get("work_experience", [])]
    except (ValidationError, TypeError):
        work = []
    try:
        education = [EducationEntry.model_validate(item).model_dump() for item in edu_raw.get("education", [])]
    except (ValidationError, TypeError):
        education = []
    try:
        skill_data = SkillOutput.model_validate(skills_raw).model_dump()
    except ValidationError:
        skill_data = SkillOutput().model_dump()

    parsed_payload: dict[str, Any] = {
        "contact_name": contact.get("name"),
        "contact_email": contact.get("email"),
        "contact_phone": contact.get("phone"),
        "contact_location": None,
        "contact_linkedin": contact.get("linkedin"),
        "contact_github": contact.get("github"),
        "contact_website": contact.get("website"),
        "contact_other": [],
        "summary_text": sections.get("summary", ""),
        "work_experience": work,
        "education": education,
        "skills_technical": skill_data.get("technical", []),
        "skills_soft": skill_data.get("soft", []),
        "skills_languages": skill_data.get("languages", []),
        "skills_certifications": skill_data.get("certifications", []),
        "projects": projects_raw.get("projects", []),
        "awards": [],
        "publications": [],
        "volunteer": [],
        "additional": {},
    }
    parsed_payload = compute_metadata(parsed_payload)

    existing = (
        await db.execute(select(ResumeParsedData).where(ResumeParsedData.resume_id == resume.id))
    ).scalar_one_or_none()
    target = existing or ResumeParsedData(resume_id=resume.id, user_id=resume.user_id, tenant_id=resume.tenant_id)
    for key, value in parsed_payload.items():
        setattr(target, key, value)
    target.parsed_at = datetime.now(UTC)
    target.parser_model = settings.WRITER_MODEL
    if existing is None:
        db.add(target)

    resume.raw_text = raw_text
    resume.original_text = raw_text
    resume.parse_status = "parsed"
    resume.parse_error = None

    await db.commit()
    await db.refresh(target)

    text_to_embed = (
        f"{target.summary_text or ''}\n"
        f"Skills: {', '.join((target.skills_all or [])[:50])}\n"
        f"Titles: {', '.join((target.primary_job_titles or [])[:3])}\n"
        f"Recent bullets: {' '.join(((target.work_experience or [{}])[0].get('bullets', [])[:5]) if target.work_experience else [])}"
    )
    chroma_client = chromadb.HttpClient(host=settings.CHROMADB_HOST, port=settings.CHROMADB_PORT)
    await store_resume_embedding(
        str(resume.id),
        text_to_embed,
        chroma_client,
        metadata={
            "user_id": str(resume.user_id),
            "resume_id": str(resume.id),
            "seniority": target.seniority_level or "",
            "domain": target.primary_domain or "",
        },
    )
    return target
